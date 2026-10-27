"""The route sheet: a process plan (capp/process_plan.py) as a document a person can review, print and sign (report 6.7, WP9).

    from capp import routesheet as RS
    plan = PP.build_plan(steps, part, context={...})
    RS.write(plan, "sheet.pdf")                # by extension: .pdf .xlsx .html .md .json
    data = RS.to_pdf_bytes(plan)               # for a download button

Per operation: number, operation and what it does, machine, tool (type, size, material), cutting parameters (speed, spindle speed,
feed, depth of cut and passes), estimated time, the feature it serves and the rule that puts it there, and flags (parameters above a
reference range, spindle load, physical-feasibility findings). Around the table: the part, where the plan came from (route source,
confidence and review flag, the before-and-after of a repair, the stated reasoning of a two-step plan), the checks, the assumptions and
the sources and status of the tables used, and a block for the signatures.

THE SHEET IS A DRAFT, and says so on every page: the tool rules and ranges are indicative (capp/process_tables.py: none is confirmed
against a handbook), the parameters are the v1 model's, there was no machining trial. When every table row used is "confirmed" the
banner changes; until then it does not. All formats show the same content (display_rows, sections), built once.

PDF needs reportlab, XLSX needs openpyxl; HTML, Markdown and JSON need nothing.
"""
import datetime
import html
import io
import json
import os

from capp import process_tables as PT

TITLE = "ROUTE SHEET"
DRAFT_TEXT = ("DRAFT FOR REVIEW. Tools and cutting parameters are indicative: the parameters are the v1 model's per-material values, the "
              "reference ranges come from web pages and the tool rules from general practice. {status} No machining trial was made. "
              "Not for the shop floor until a process engineer or the supervisor has reviewed and signed it.")
CONFIRMED_TEXT = "All reference ranges used are confirmed against a handbook (see Sources). Tool rules are still indicative."
COLUMNS = (("seq", "No."), ("operation", "Operation"), ("machine", "Machine"), ("tool", "Tool"), ("vc", "Vc m/min"), ("rpm", "n rpm"),
           ("feed", "Feed"), ("depth", "ap mm x passes"), ("time", "Time min"), ("why", "Feature / rule"), ("notes", "Notes"))
WIDTHS_MM = (9, 44, 15, 46, 13, 12, 25, 18, 13, 40, 42)         # for PDF landscape A4 (277 mm usable) and the XLSX column widths


class RouteSheetError(RuntimeError):
    pass


# ---------------------------------------------------------------- the content, once
def _f(x, nd=1, dash="-"):
    return dash if x is None else f"{x:.{nd}f}"


def draft_banner(plan) -> str:
    st = plan["tables"]["status"]
    n = plan["tables"]["ranges_used"]
    if n and st["confirmed"] == n:
        return DRAFT_TEXT.format(status=CONFIRMED_TEXT)
    parts = [f"{st[k]} {k}" for k in ("confirmed", "web", "mixed", "indicative") if st[k]]
    text = f"Of the {n} reference ranges used, " + ", ".join(parts) + "; none is confirmed against a handbook." if n and not st["confirmed"] else \
        f"Of the {n} reference ranges used, " + ", ".join(parts) + "." if n else "No reference range was used."
    return DRAFT_TEXT.format(status=text)


def header_items(plan, generated_at) -> list:
    p, c = plan["part"], plan["context"]
    dims = p["dims"]
    dim_text = "reference part of v1 (Ø25, L50, depth 20): no dimensions given" if dims is None else \
        f"Ø{dims['diameter']:g} x L{dims['length']:g}, depth {dims['depth']:g}, width {dims['width']:g} mm" + \
        (f", hole Ø{dims['hole_diameter']:g}" if dims.get("hole_diameter") else ", no hole diameter given")
    items = [("Part", p["part_id"] or "-"), ("Material", p["material"]), ("Features", ", ".join(p["features"]) or "-"),
             ("Tolerance", p["tolerance"] or "-"), ("Batch size", str(p["batch_size"])), ("Dimensions", dim_text),
             ("Required Ra", "-" if not p["required_ra_um"] else f"{p['required_ra_um']:g} µm"),
             ("Machines", plan["machine"]["name"] + f" ({plan['machine']['spindle_power_kw']:g} kW spindle, assumed)")]
    src = c.get("source") or c.get("provenance")
    if src:
        items.append(("Route source", str(src).replace("_", " ") + (f" ({c['system']})" if c.get("system") else "")))
    if c.get("confidence") is not None:
        flag = "FLAGGED for human review" if c.get("flagged") else "not flagged"
        thr = f", threshold {c['confidence_flag']:g}" if c.get("confidence_flag") is not None else ""
        items.append(("Plan confidence", f"{c['confidence']:.2f} ({flag}{thr})"))
    items.append(("Generated", generated_at))
    return items


def display_rows(plan) -> list:
    """The table, as display strings: [{"kind": "op" | "changeover", ...}] in route order."""
    rows = []
    for e in plan["operations"]:
        if e["changeover_before"]:
            co = next((c for c in plan["changeovers"] if c["before_seq"] == e["seq"]), None)
            frm = co["from"] if co else "?"
            rows.append({"kind": "changeover", "text": f"MACHINE CHANGEOVER: {frm} -> {e['machine']} "
                         f"({e['time']['changeover_min']:g} min, counted in the next operation's time)"})
        t, p = e["tool"], e["params"]
        tool = "-" if t is None else f"{t['type']}" + (f", {t['size']}{'*' if t['size_assumed'] else ''}" if t["size"] else "") + \
            ("" if t["material"] == "-" else f"; {t['material']}")
        if p:
            feed = f"{p['feed']:g} {p['feed_unit']}" + ("" if p["feed_rate_mm_min"] is None else f" ({p['feed_rate_mm_min']:.0f} mm/min)")
            depth = "-" if p["ap_mm"] is None else f"{p['ap_mm']:.2f} x {p['passes']}"
            load = "" if p["load"] is None else f"load {p['load'] * 100:.0f}%"
        else:
            feed = depth = load = "-"
        notes = list(e["flags"]) + ([load] if load and p and p["load"] and p["load"] > 0.5 else [])
        rows.append({"kind": "op", "seq": str(e["seq"]), "operation": e["op"], "description": e["description"],
                     "machine": "-" if e["op"] == "Inspection" else e["machine"] + ("*" if e["machine_assumed"] else ""), "tool": tool,
                     "vc": "-" if not p else f"{p['vc_m_min']:g}", "rpm": "-" if not p else f"{p['rpm']:.0f}", "feed": feed, "depth": depth,
                     "time": f"{e['time_min']:.2f}", "why": e["why"] + "".join(f" [{r}]" for r in e["rules"] if not r.startswith("ORD-EDGE")),
                     "notes": "; ".join(notes)})
    return rows


def totals_text(plan) -> list:
    t = plan["totals"]
    lines = [f"Operations: {t['n_operations']}, machine changeovers: {t['n_changeovers']}",
             f"Estimated time per part: {t['total_min']:.2f} min = cutting {t['cutting_min']:.2f} + tool changes {t['tool_change_min']:.2f} + "
             f"position changes {t['position_change_min']:.2f} + changeovers {t['changeover_min']:.2f}",
             f"Batch of {plan['part']['batch_size']}: {t['per_batch_hours']:.1f} h of machine time (without set-up of the batch, waiting and inspection of the batch)"]
    ph = plan["physics"]
    lines.append(f"Cutting energy {ph['cutting_energy_kwh']:.2f} kWh per part; peak spindle load {ph['max_utilisation'] * 100:.0f}% of the assumed machine")
    return lines


def tables_used(plan) -> tuple:
    """(tool rows, range rows) of the tables this plan used, in the shape of PT.rows(): the rows a reviewer has to check."""
    mat = plan["part"]["material"]
    tools, ranges, seen_t, seen_r = [], [], set(), set()
    for e in plan["operations"]:
        if not e["known"] or e["tool"] is None:
            continue
        key = f"{e['op']} on {e['machine']} | {mat}" if e["tool"].get("machine_override") else f"{e['op']} | {mat}"
        if key not in seen_t:
            seen_t.add(key)
            tools.append({"table": "tool", "key": key, "value": f"{e['tool']['type']}; {e['tool']['material']}", "status": "indicative",
                          "sources": "", "note": "general machining practice"})
        rc = e.get("range_check")
        if rc and rc["range"] is not None:
            r = rc["range"]
            rkey = f"{r['speed_class']} | {mat}"
            if rkey not in seen_r:
                seen_r.add(rkey)
                unit = "mm/tooth" if r["speed_class"] == "mill" else "mm/rev"
                ranges.append({"table": "range", "key": rkey, "status": r["status"], "sources": ",".join(r["sources"]),
                               "value": f"Vc {r['vc'][0]:g}-{r['vc'][1]:g} m/min; feed {r['feed'][0]:g}-{r['feed'][1]:g} {unit}",
                               "note": r["note"] + (" " + r["source_note"] if r.get("source_note") else "")})
    return tools, ranges


def sections(plan, generated_at=None) -> list:
    """The text blocks after the table: [(title, [lines])]. Empty blocks are left out."""
    out = [("Totals", totals_text(plan))]
    v, fz, ph = plan["verifier"], plan["feasibility"], plan["physics"]
    checks = ["Rule check of the route (precedence, coverage, vocabulary): " + ("PASSED" if v["valid"] else "FAILED: " + "; ".join(v["messages"]))]
    if fz["error"]:
        checks.append(fz["error"])
    checks += [f"Physical feasibility, {'hard finding' if f['kind'] == 'infeasible' else 'risk'} {f['rule_id']}: {f['message']}" for f in fz["findings"]]
    if not fz["error"] and not fz["hard"] and not fz["risks"]:
        checks.append("Physical feasibility: no hard finding and no risk among the rules that could be evaluated for this part")
    checks += [f"Not evaluated: {rid} ({why})" for rid, why in fz["not_evaluated"]]
    checks += [f"Physics: {m}" for m in ph["flags"]]
    out.append(("Checks", checks))
    c = plan["context"]
    if c.get("selection") or c.get("n_samples"):
        out.append(("How the route was chosen", [f"{k}: {c[k]}" for k in ("system", "selection", "n_samples") if c.get(k)]))
    rep = plan.get("repair")
    if rep:
        lines = ["The planner's raw route and the route of this sheet are " + ("identical." if rep["identical"] else "different:")]
        if not rep["identical"]:
            lines.append("raw route: " + " > ".join(rep["raw"]))
            lines.append("this sheet: " + " > ".join(f"{op}{'+' if m == 'inserted' else ''}" for op, m in zip(rep["final"], rep["marks"]))
                         + "   (+ = inserted by the repair)")
            if rep["removed"]:
                lines.append("removed: " + ", ".join(rep["removed"]))
            if rep["reordered"]:
                lines.append("the order of the kept operations was changed")
        out.append(("Before and after repair", lines))
    if c.get("reasoning"):
        out.append(("Stated reasoning of the planner", [str(c["reasoning"]).strip()]))
    out.append(("Assumptions", plan["assumptions"]))
    _, used = tables_used(plan)
    src = [f"reference range {r['key']}: {r['value']} (status {r['status']}, source {r['sources'] or '-'})" for r in used]
    src.append("tool types and tool materials: rule table of general machining practice, status 'indicative'")
    src += [f"[{k}] {PT.SOURCES[k][0]}: {PT.SOURCES[k][1]}" for k in plan["tables"]["sources"]]
    out.append(("Sources and status of the tables", src))
    return [(t, [x for x in lines if x]) for t, lines in out if lines]


def _now():
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


# ---------------------------------------------------------------- Markdown and JSON
def to_markdown(plan, generated_at=None) -> str:
    g = generated_at or _now()
    L = [f"# {TITLE}", "", f"> **{draft_banner(plan)}**", ""]
    L += [f"- **{k}:** {v}" for k, v in header_items(plan, g)]
    L += ["", "| " + " | ".join(h for _, h in COLUMNS) + " |", "|" + "---|" * len(COLUMNS)]
    for r in display_rows(plan):
        if r["kind"] == "changeover":
            L.append(f"| | *{r['text']}* |" + " |" * (len(COLUMNS) - 2))
        else:
            cells = [r["seq"], f"**{r['operation']}**<br>{r['description']}"] + [r[k].replace("|", "/") for k, _ in COLUMNS[2:]]
            L.append("| " + " | ".join(cells) + " |")
    L.append("")
    for title, lines in sections(plan, g):
        L += [f"## {title}", ""] + [f"- {x}" if len(lines) > 1 else x for x in lines] + [""]
    L += ["## Approval", "", "Prepared by: ____________  Checked by: ____________  Approved by: ____________  Date: ____________"]
    return "\n".join(L) + "\n"


def to_json(plan, generated_at=None) -> str:
    return json.dumps(dict(plan, generated_at=generated_at or _now(), banner=draft_banner(plan)), indent=2, ensure_ascii=False, default=str)


# ---------------------------------------------------------------- HTML
_CSS = """
body{font:12px/1.4 Arial,Helvetica,sans-serif;color:#1a1a1a;margin:18px;background:#fff}
h1{font-size:20px;margin:0 0 6px}h2{font-size:14px;margin:18px 0 4px;border-bottom:1px solid #bbb}
.banner{border:2px solid #b45309;background:#fff7ed;color:#7c2d12;padding:8px 10px;margin:8px 0;font-weight:bold}
table{border-collapse:collapse;width:100%}th,td{border:1px solid #999;padding:3px 5px;vertical-align:top;text-align:left}
th{background:#e5e7eb}tr.co td{background:#fef3c7;font-style:italic}td.n{text-align:right;white-space:nowrap}
.hdr td{border:none;padding:1px 10px 1px 0}.hdr td:first-child{font-weight:bold;white-space:nowrap}
.note{color:#7c2d12}.sign{margin-top:26px}.sign span{display:inline-block;width:30%}
@media print{body{margin:6mm}@page{size:A4 landscape;margin:8mm}}
"""


def to_html(plan, generated_at=None) -> str:
    g = generated_at or _now()
    e = html.escape
    out = ["<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>",
           f"<title>{TITLE} {e(plan['part']['part_id'] or '')}</title><style>{_CSS}</style></head><body>",
           f"<h1>{TITLE}</h1><div class='banner'>{e(draft_banner(plan))}</div><table class='hdr'>"]
    out += [f"<tr><td>{e(k)}</td><td>{e(str(v))}</td></tr>" for k, v in header_items(plan, g)]
    out.append("</table><h2>Operations</h2><table><thead><tr>" + "".join(f"<th>{e(h)}</th>" for _, h in COLUMNS) + "</tr></thead><tbody>")
    num = {"vc", "rpm", "time", "seq"}
    for r in display_rows(plan):
        if r["kind"] == "changeover":
            out.append(f"<tr class='co'><td colspan='{len(COLUMNS)}'>{e(r['text'])}</td></tr>")
            continue
        cells = []
        for k, _ in COLUMNS:
            if k == "operation":
                cells.append(f"<td><b>{e(r['operation'])}</b><br>{e(r['description'])}</td>")
            elif k == "notes":
                cells.append(f"<td class='note'>{e(r['notes'])}</td>")
            else:
                cells.append(f"<td{' class=n' if k in num else ''}>{e(r[k])}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</tbody></table><p>* assumed (see Assumptions).</p>")
    for title, lines in sections(plan, g):
        out.append(f"<h2>{e(title)}</h2>")
        if len(lines) == 1:
            out.append(f"<p style='white-space:pre-wrap'>{e(lines[0])}</p>")
        else:
            out.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in lines) + "</ul>")
    out.append("<div class='sign'><span>Prepared by: ____________</span><span>Checked by: ____________</span><span>Approved by: ____________</span></div>"
               "<p>Date: ____________</p></body></html>")
    return "".join(out)


# ---------------------------------------------------------------- XLSX
def to_xlsx(plan, path_or_file, generated_at=None):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RouteSheetError("the XLSX route sheet needs openpyxl: pip install openpyxl") from exc
    g = generated_at or _now()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Route sheet"
    n = len(COLUMNS)
    ws.append([TITLE])
    ws["A1"].font = Font(bold=True, size=16)
    ws.append([draft_banner(plan)])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws["A2"].font = Font(bold=True, color="7C2D12")
    ws["A2"].fill = PatternFill("solid", fgColor="FFF7ED")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 48
    for k, v in header_items(plan, g):
        ws.append([k, str(v)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=n)
    ws.append([])
    ws.append([h for _, h in COLUMNS])
    head_row = ws.max_row
    for c in ws[head_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E5E7EB")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    numeric = {"seq": int, "vc": float, "rpm": float, "time": float}
    for r in display_rows(plan):
        if r["kind"] == "changeover":
            ws.append([r["text"]])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=n)
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="FEF3C7")
            ws.cell(ws.max_row, 1).font = Font(italic=True)
            continue
        cells = []
        for k, _ in COLUMNS:
            v = r["operation"] + "\n" + r["description"] if k == "operation" else r[k]
            if k in numeric:
                try:
                    v = numeric[k](v)
                except ValueError:
                    pass
            cells.append(v)
        ws.append(cells)
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(WIDTHS_MM, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, w * 0.95)
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{head_row}:{head_row}"
    ns = wb.create_sheet("Checks and notes")
    for title, lines in sections(plan, g):
        ns.append([title])
        ns.cell(ns.max_row, 1).font = Font(bold=True)
        for x in lines:
            ns.append([x])
            ns.cell(ns.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ns.append([])
    ns.column_dimensions["A"].width = 150
    st = wb.create_sheet("Tables used")
    st.append(["table", "key", "value", "status", "sources", "note"])
    for c in st[1]:
        c.font = Font(bold=True)
    tools, ranges = tables_used(plan)
    for r in tools + ranges:
        st.append([r["table"], r["key"], r["value"], r["status"], r["sources"], r["note"]])
    for col, w in zip("ABCDEF", (8, 36, 80, 12, 10, 70)):
        st.column_dimensions[col].width = w
    st.append([])
    for k in plan["tables"]["sources"]:
        st.append(["source", k, f"{PT.SOURCES[k][0]}  {PT.SOURCES[k][1]}"])
    wb.save(path_or_file)
    return path_or_file


# ---------------------------------------------------------------- PDF
_PDF_SUBST = {"→": "->", "≥": ">=", "≤": "<=", "−": "-", "–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "₹": "Rs", "…": "...",
              "\u00a0": " ",
              "µ": "u"}          # the micro sign is in WinAnsi, but some viewers draw nothing for it with the unembedded base fonts


def pdf_text(s) -> str:
    """Text the standard PDF fonts can draw (Latin-1 with Ø and °; µ is written u); anything else becomes '?', never an error."""
    s = "" if s is None else str(s)
    for a, b in _PDF_SUBST.items():
        s = s.replace(a, b)
    return s.encode("cp1252", "replace").decode("cp1252")


def to_pdf(plan, path_or_file, generated_at=None):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RouteSheetError("the PDF route sheet needs reportlab: pip install reportlab (or use the XLSX or HTML sheet)") from exc
    from xml.sax.saxutils import escape
    g = generated_at or _now()
    base = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=base["BodyText"], fontSize=7, leading=8.4)
    head = ParagraphStyle("head", parent=small, fontName="Helvetica-Bold")
    body = ParagraphStyle("body", parent=base["BodyText"], fontSize=8.5, leading=10.5)
    h2 = ParagraphStyle("h2", parent=base["Heading3"], fontSize=10.5, spaceBefore=8, spaceAfter=2)
    banner = ParagraphStyle("banner", parent=body, fontName="Helvetica-Bold", textColor=colors.HexColor("#7c2d12"), backColor=colors.HexColor("#fff7ed"),
                            borderColor=colors.HexColor("#b45309"), borderWidth=1.2, borderPadding=5, spaceAfter=8, spaceBefore=4)

    def P(text, style=small):
        return Paragraph(escape(pdf_text(text)).replace("\n", "<br/>"), style)

    story = [Paragraph(TITLE, base["Title"]), Paragraph(escape(pdf_text(draft_banner(plan))), banner)]
    hdr = [[P(k, head), P(v, body)] for k, v in header_items(plan, g)]
    ht = Table(hdr, colWidths=[32 * mm, 245 * mm])
    ht.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 1), ("TOPPADDING", (0, 0), (-1, -1), 1)]))
    story += [ht, Spacer(1, 6), Paragraph("Operations", h2)]
    data = [[P(h, head) for _, h in COLUMNS]]
    spans, shade = [], []
    for r in display_rows(plan):
        i = len(data)
        if r["kind"] == "changeover":
            data.append([P(r["text"], head)] + [""] * (len(COLUMNS) - 1))
            spans.append(("SPAN", (0, i), (-1, i)))
            shade.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#fef3c7")))
            continue
        row = []
        for k, _ in COLUMNS:
            if k == "operation":
                row.append(Paragraph(f"<b>{escape(pdf_text(r['operation']))}</b><br/>{escape(pdf_text(r['description']))}", small))
            elif k == "notes":
                row.append(Paragraph(f"<font color='#7c2d12'>{escape(pdf_text(r['notes']))}</font>", small))
            else:
                row.append(P(r[k]))
        data.append(row)
    t = Table(data, colWidths=[w * mm for w in WIDTHS_MM], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5e7eb")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 2.5), ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                           ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)] + spans + shade))
    story += [t, P("* assumed (see Assumptions).", small)]
    for title, lines in sections(plan, g):
        block = [Paragraph(escape(pdf_text(title)), h2)]
        block += [Paragraph(("" if len(lines) == 1 else "&bull; ") + escape(pdf_text(x)).replace("\n", "<br/>"), body) for x in lines]
        story.append(KeepTogether(block) if len(lines) < 8 else block[0])
        if len(lines) >= 8:
            story += block[1:]
    sig = Table([[P("Prepared by: ______________", body), P("Checked by: ______________", body), P("Approved by: ______________", body),
                  P("Date: ____________", body)]], colWidths=[70 * mm] * 4)
    story += [Spacer(1, 14), sig]
    label = pdf_text(f"{TITLE} {plan['part']['part_id'] or ''} - DRAFT".strip())

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#7c2d12"))
        canvas.drawString(10 * mm, 6 * mm, label + " - indicative tools and parameters, see the banner on page 1")
        canvas.drawRightString(287 * mm, 6 * mm, f"page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path_or_file, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
                            title=pdf_text(f"Route sheet {plan['part']['part_id'] or ''}".strip()), author="LLM-CAPP (MTP-1)")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path_or_file


def to_pdf_bytes(plan, generated_at=None) -> bytes:
    buf = io.BytesIO()
    to_pdf(plan, buf, generated_at)
    return buf.getvalue()


def to_xlsx_bytes(plan, generated_at=None) -> bytes:
    buf = io.BytesIO()
    to_xlsx(plan, buf, generated_at)
    return buf.getvalue()


# ---------------------------------------------------------------- by extension
WRITERS = {".pdf": to_pdf, ".xlsx": to_xlsx}
TEXT_WRITERS = {".html": to_html, ".htm": to_html, ".md": to_markdown, ".json": to_json}


def write(plan, path: str, generated_at: str = None) -> str:
    """Write the sheet in the format of the file extension (.pdf .xlsx .html .md .json)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in WRITERS:
        WRITERS[ext](plan, path, generated_at)
    elif ext in TEXT_WRITERS:
        with open(path, "w", encoding="utf-8") as f:
            f.write(TEXT_WRITERS[ext](plan, generated_at))
    else:
        raise RouteSheetError(f"unknown format {ext!r}: use one of {', '.join(sorted(list(WRITERS) + list(TEXT_WRITERS)))}")
    return path
