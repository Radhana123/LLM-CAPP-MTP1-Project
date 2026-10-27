"""Expert benchmark pipeline (MTP-1 WP2, Step 16): from a list of real parts to a sealed benchmark with expert reference plans, a
human-human agreement figure (M10) and a datasheet.

The plans themselves are written by people, so this module is the tooling around them:

    1. spec        the parts: material, features, tolerance, batch, REAL dimensions (hole diameter optional), source of the part.
                   write_spec_template() gives an Excel/CSV template; read_spec() reads and validates it (row-wise error messages);
                   stratified_pick() can draw a placeholder spec from a synthetic set to rehearse the whole chain.
    2. sheets      write_annotation_workbook(): one Excel workbook per annotator, a row per part and alternative, one dropdown cell
                   per step (the 41 operations), no plan pre-filled (a pre-filled plan would anchor the annotator, and a plan
                   pre-filled by the Route Builder would favour the Route Builder in every later comparison). The primary workbook
                   has review columns for the supervisor; the second annotator gets a subset (second_subset()).
    3. import      read_annotations() reads a filled workbook (unknown operation names are errors, with suggestions);
                   build_benchmark() checks every plan against the rule base (precedence graph, missing/invented operations)
                   and the physical-feasibility rules with the real dimensions, and builds the records of the benchmark dataset.
    4. agreement   agreement(): M10, the M4 measures (exact match, LCS ratio, Kendall tau) between the second person's plan and the
                   reference plan(s), on the parts both did, with a bootstrap interval. It is the ceiling for M4.
    5. datasheet   datasheet(): a markdown datasheet generated from the data (composition, plans, review, conflicts, M10, limits).

A reference plan that the rule base calls invalid is KEPT as the reference (the expert is the ground truth of the benchmark) and listed
as a conflict: either the expert or the rule base is wrong, and which one is a finding for the supervisor, not for the code.

Records follow capp/dataset.py (split "bench", dims may have hole_diameter, reference_routes, second_annotator_routes, reference_meta).
Everything except the Excel/CSV/JSON reading and writing is a pure function.
"""
import csv
import difflib
import json
import os
import random
import re
import statistics
from collections import Counter, defaultdict

from capp import legacy  # noqa: F401
from capp import dataset as D
from capp import feasibility as F
from capp import metrics as M
from capp import summary as S
from capp import verify as V
import feature_vocab as fv          # noqa: E402  (on the path through capp.dataset)
import route_builder as rb          # noqa: E402
from material_tokens import MATERIALS, TOLERANCE_CATEGORIES      # noqa: E402

SPEC_COLUMNS = ("part_id", "name", "source", "drawing", "material", "features", "tolerance", "batch_size",
                "diameter", "length", "depth", "width", "hole_diameter", "required_ra_um", "notes")
REQUIRED_COLUMNS = ("name", "source", "material", "features", "tolerance", "batch_size", "diameter", "length", "depth", "width")
MAX_STEPS = 24
ALT_LABELS = ("A", "B", "C")
REVIEW_STATUSES = ("accepted", "changed", "rejected")
MATERIAL_ALIASES = {"aluminium": "Aluminum", "al": "Aluminum", "castiron": "Cast Iron", "gci": "Cast Iron"}
OPERATIONS = sorted(V.CANONICAL_OPS)
BOOTSTRAP_RESAMPLES = 2000


# ---------------------------------------------------------------- small helpers
def _key(s) -> str:
    return re.sub(r"[\s_\-/]+", "", str(s)).lower()


def _match(value, choices, aliases=None):
    """The canonical choice that `value` names (case, spaces, underscores and '/' ignored), or None."""
    k = _key(value)
    for c in choices:
        if _key(c) == k:
            return c
    if aliases and k in aliases:
        return aliases[k]
    return None


def _suggest(value, choices, n=3):
    keys = {_key(c): c for c in choices}
    return [keys[k] for k in difflib.get_close_matches(_key(value), list(keys), n=n, cutoff=0.6)]


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _num(v):
    if _blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError("not a number")
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).strip().replace(",", "."))


def _text(v) -> str:
    return "" if _blank(v) else str(v).strip()


# ---------------------------------------------------------------- 1. spec
def parse_part_row(row: dict, index: int):
    """(part or None, errors). `row` has the SPEC_COLUMNS as keys (missing keys count as blank). `index` numbers the generated ids."""
    errs = []
    for c in REQUIRED_COLUMNS:
        if _blank(row.get(c)):
            errs.append(f"'{c}' is empty")
    pid = _text(row.get("part_id")) or f"BENCH_{index:03d}"

    material = None
    if not _blank(row.get("material")):
        material = _match(row["material"], MATERIALS, MATERIAL_ALIASES)
        if material is None:
            errs.append(f"material {row['material']!r} is not one of {', '.join(MATERIALS)}")

    feats = []
    if not _blank(row.get("features")):
        for raw in re.split(r"[;,]", str(row["features"])):
            if _blank(raw):
                continue
            f = _match(raw, fv.get_all_features())
            if f is None:
                sug = _suggest(raw, fv.get_all_features())
                errs.append(f"unknown feature {raw.strip()!r}" + (f" (did you mean {', '.join(sug)}?)" if sug else ""))
            elif f not in feats:
                feats.append(f)

    tol = None
    if not _blank(row.get("tolerance")):
        m = re.fullmatch(r"\s*([0-9]*[.,]?[0-9]+)\s*(mm)?\s*", str(row["tolerance"]), re.IGNORECASE)
        tol = f"{float(m.group(1).replace(',', '.')):g}mm" if m else None
        if tol not in TOLERANCE_CATEGORIES:
            errs.append(f"tolerance {row['tolerance']!r} is not one of {', '.join(TOLERANCE_CATEGORIES)}")
            tol = None

    batch = None
    if not _blank(row.get("batch_size")):
        try:
            b = _num(row["batch_size"])
            if b != int(b):
                raise ValueError
            batch = int(b)
            if not D.BATCH_LIMITS[0] <= batch <= D.BATCH_LIMITS[1]:
                errs.append(f"batch_size {batch} outside {D.BATCH_LIMITS}")
        except ValueError:
            errs.append(f"batch_size {row['batch_size']!r} is not a whole number")

    dims = {}
    for k in D.DIM_KEYS:
        if _blank(row.get(k)):
            continue
        try:
            v = _num(row[k])
            lo, hi = D.DIM_LIMITS[k]
            if not lo <= v <= hi:
                errs.append(f"{k} {v:g} mm outside [{lo:g}, {hi:g}]")
            dims[k] = v
        except ValueError:
            errs.append(f"{k} {row[k]!r} is not a number")
    if not _blank(row.get("hole_diameter")):
        try:
            v = _num(row["hole_diameter"])
            if not D.HOLE_LIMITS[0] <= v <= D.HOLE_LIMITS[1]:
                errs.append(f"hole_diameter {v:g} mm outside {D.HOLE_LIMITS}")
            dims["hole_diameter"] = v
        except ValueError:
            errs.append(f"hole_diameter {row['hole_diameter']!r} is not a number")

    ra = None
    if not _blank(row.get("required_ra_um")):
        try:
            ra = _num(row["required_ra_um"])
            if not D.BENCH_RA_LIMITS[0] <= ra <= D.BENCH_RA_LIMITS[1]:
                errs.append(f"required_ra_um {ra:g} outside {D.BENCH_RA_LIMITS}")
        except ValueError:
            errs.append(f"required_ra_um {row['required_ra_um']!r} is not a number")
    if errs:
        return None, errs
    part = {"part_id": pid, "split": "bench", "archetype": _text(row["name"]), "in_distribution": None, "material": material,
            "features": feats, "tolerance": tol, "batch_size": batch,
            "dims": {k: (int(v) if float(v).is_integer() else v) for k, v in dims.items()},
            "required_ra_um": ra, "machine_preference": "auto"}
    part.update(D.derived_fields(part))
    part.update({"name": _text(row["name"]), "source": _text(row["source"]), "drawing": _text(row.get("drawing")),
                 "notes": _text(row.get("notes"))})
    return part, []


def read_table(path: str) -> list:
    """The rows of a .csv, .xlsx or .json (list of part records) as dicts keyed by lower-case header. Empty rows are skipped."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["Parts"] if "Parts" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = [list(r) for r in csv.reader(f)]
    if not rows:
        return []
    header = [re.sub(r"\s+", "_", _text(h).lower()) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(_blank(c) for c in r):
            continue
        out.append({h: c for h, c in zip(header, r) if h})
    return out


def read_spec(path: str):
    """(parts, errors). errors are 'row N: ...' strings (N = the row number in the file, header = 1). A .json spec (written by
    scripts/benchmark.py sheets) holds finished records and is validated as they are."""
    if os.path.splitext(path)[1].lower() == ".json":
        parts = read_table(path)
        errs = []
        for p in parts:
            for e in D.validate_part(dict(p, reference_routes=p.get("reference_routes") or [["_"]])):
                errs.append(f"{p.get('part_id')}: {e}")
        return parts, errs
    rows = read_table(path)
    parts, errors, seen = [], [], Counter()
    for i, row in enumerate(rows, 1):
        part, errs = parse_part_row(row, i)
        errors += [f"row {i + 1}: {e}" for e in errs]
        if part:
            parts.append(part)
            seen[part["part_id"]] += 1
    errors += [f"part_id {pid} appears {n} times" for pid, n in seen.items() if n > 1]
    if not rows:
        errors.append("the spec has no part rows")
    return parts, errors


def _op_features() -> dict:
    out = defaultdict(set)
    for f, info in fv.FEATURE_TO_OPERATIONS.items():
        for alt in info["alternatives"]:
            for op in alt:
                out[op].add(f)
    return out


def machine_class(op: str) -> str:
    return "lathe" if op in rb.LATHE_ONLY_OPS else "milling" if op in rb.MILLING_ONLY_OPS else "either"


def _style_header(ws, widths=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE6F0")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def write_spec_template(path: str):
    """An empty spec: Excel (sheets Parts, Lists, README) or, for a .csv path, just the header row."""
    if path.lower().endswith(".csv"):
        with open(path, "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerow(SPEC_COLUMNS)
        return path
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append(list(SPEC_COLUMNS))
    _style_header(ws, {"A": 11, "B": 28, "C": 30, "D": 22, "E": 12, "F": 34, "G": 10, "H": 10, "I": 10, "J": 10, "K": 9, "L": 9, "M": 13, "N": 13, "O": 30})
    lists = wb.create_sheet("Lists")
    lists.append(["material", "tolerance", "feature", "machine", "usual operations (alternatives)"])
    feats = fv.get_all_features()
    for i in range(max(len(MATERIALS), len(TOLERANCE_CATEGORIES), len(feats))):
        lists.append([list(MATERIALS)[i] if i < len(MATERIALS) else None,
                      list(TOLERANCE_CATEGORIES)[i] if i < len(TOLERANCE_CATEGORIES) else None,
                      feats[i] if i < len(feats) else None,
                      fv.get_machine_type(feats[i]) if i < len(feats) else None,
                      " | ".join(" > ".join(a) for a in fv.FEATURE_TO_OPERATIONS[feats[i]]["alternatives"]) if i < len(feats) else None])
    _style_header(lists, {"A": 12, "B": 10, "C": 18, "D": 9, "E": 90})
    for col, n in (("E", len(MATERIALS)), ("G", len(TOLERANCE_CATEGORIES))):
        src = "Lists!$A$2:$A$%d" % (n + 1) if col == "E" else "Lists!$B$2:$B$%d" % (n + 1)
        dv = DataValidation(type="list", formula1="=" + src, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}300")
    readme = wb.create_sheet("README")
    for line in SPEC_README:
        readme.append([line])
    readme.column_dimensions["A"].width = 130
    wb.save(path)
    return path


SPEC_README = [
    "Expert benchmark: parts specification (one row per part, sheet 'Parts').",
    "",
    "name, source: required. The source is where the part comes from (textbook and exercise number, handbook example, department drawing id).",
    "material: one of the list in sheet 'Lists'. tolerance: the tightest tolerance of the part, one of the list (e.g. 0.02mm).",
    "features: the machining features of the part, separated by ';', names as in sheet 'Lists' (e.g. Hole; Thread; Chamfer).",
    "batch_size: pieces per batch (1 to 10000).",
    "diameter, length, depth, width: REAL dimensions in mm of the stock / part envelope (the same four numbers the app asks for).",
    "hole_diameter: mm, optional but wanted: the main hole diameter if the part has a hole (needed by the deep-hole and power checks).",
    "required_ra_um: surface roughness Ra from the drawing in micrometres, leave empty if the drawing gives none.",
    "drawing: file name or reference of the drawing, if there is one. notes: anything the annotator should know.",
    "part_id: leave empty to get BENCH_001, BENCH_002 ... in row order.",
    "",
    "Do not change the header row. Do not add the plans here: they are written by the annotators in their own workbooks.",
]


def stratified_pick(parts: list, n: int, seed: int = 7) -> list:
    """n parts spread over (machine profile, complexity) as evenly as the data allows, deterministic for a seed, in the original order."""
    rng = random.Random(seed)
    groups = defaultdict(list)
    for i, p in enumerate(parts):
        groups[(p["machine_profile"], S.complexity(len(p["features"])))].append(i)
    for idx in groups.values():
        rng.shuffle(idx)
    picked = []
    keys = sorted(groups)
    while len(picked) < min(n, len(parts)):
        for k in keys:
            if groups[k] and len(picked) < n:
                picked.append(groups[k].pop())
    return [parts[i] for i in sorted(picked)]


def second_subset(parts: list, n: int, seed: int = 7) -> list:
    """The parts the second annotator plans (15 to 20 in the protocol): a stratified pick of the benchmark."""
    return stratified_pick(parts, n, seed)


def spec_from_dataset(parts: list, source: str) -> list:
    """Rows of a PLACEHOLDER spec from synthetic parts (to rehearse the pipeline): new ids, the original id in the notes."""
    rows = []
    for i, p in enumerate(parts, 1):
        row = {"part_id": f"BENCH_{i:03d}", "name": p["archetype"], "source": source, "drawing": "",
               "material": p["material"], "features": ";".join(p["features"]), "tolerance": p["tolerance"],
               "batch_size": p["batch_size"], "required_ra_um": p.get("required_ra_um"),
               "notes": f"synthetic part {p['part_id']}: placeholder, not a real drawing"}
        row.update(p["dims"])
        rows.append(row)
    return rows


def write_spec_rows(rows: list, path: str):
    if path.lower().endswith(".xlsx"):
        write_spec_template(path)
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb["Parts"]
        for r in rows:
            ws.append([r.get(c) for c in SPEC_COLUMNS])
        wb.save(path)
    else:
        with open(path, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(SPEC_COLUMNS)
            for r in rows:
                w.writerow([r.get(c, "") if r.get(c) is not None else "" for c in SPEC_COLUMNS])
    return path


# ---------------------------------------------------------------- 2. annotation workbooks
def part_summary(p: dict) -> str:
    d = p["dims"]
    dims = f"Ø{d['diameter']:g} x L{d['length']:g}, depth {d['depth']:g}, width {d['width']:g} mm" + (
        f", hole Ø{d['hole_diameter']:g}" if "hole_diameter" in d else "")
    ra = f", Ra {p['required_ra_um']:g}" if p.get("required_ra_um") else ""
    return (f"{p['name']} | {p['material']} | {'; '.join(p['features'])} | tol {p['tolerance']}{ra} | batch {p['batch_size']} | {dims}")


ANNOTATION_README = [
    "Process plan annotation: write the machining route (the order of operations) for each part.",
    "",
    "1. Sheet 'Routes': for every part, row A is your plan. Write one operation per cell, left to right, in the order you would run them.",
    "   Pick the operation from the dropdown (names are fixed: sheet 'Operations' lists the 41 with the machine they need and the",
    "   features they usually serve). A cell that is not in the list is rejected at import.",
    "2. End every plan with Inspection (the rule base expects it last).",
    "3. Rows B and C: only if another plan is EQUALLY acceptable to you (another order of independent operations, another process chain).",
    "   Leave them empty otherwise.",
    "4. Sheet 'Parts' has the part data (material, features, tolerance, batch, real dimensions). Read it, then plan from it.",
    "5. Plan independently. Do not look at another annotator's plan or at any program output for these parts: the agreement between",
    "   two independent people is a measurement, and it is spoiled if one of them has seen the other's plan.",
    "6. Machine changeover markers are not written; only operations.",
    "7. Comments (a doubt, a reason for an alternative) go in the 'comment' column of the Routes sheet.",
]
REVIEW_README = [
    "",
    "Supervisor review (primary workbook only): in sheet 'Parts', per part, set review_status to accepted / changed / rejected,",
    "write your name in reviewed_by and, if you changed or rejected the plan, edit the plan in sheet 'Routes' and say why in review_comment.",
]


def write_annotation_workbook(parts: list, path: str, annotator: str, role: str = "primary", alternatives: int = 2,
                              max_steps: int = MAX_STEPS) -> str:
    """One workbook for one annotator. role is 'primary' (reference plans, with review columns) or 'second' (independent plan)."""
    if role not in ("primary", "second"):
        raise ValueError("role must be 'primary' or 'second'")
    if not 1 <= alternatives <= len(ALT_LABELS):
        raise ValueError(f"alternatives must be 1..{len(ALT_LABELS)}")
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    for line in ANNOTATION_README + (REVIEW_README if role == "primary" else []):
        readme.append([line])
    readme.column_dimensions["A"].width = 130

    ps = wb.create_sheet("Parts")
    cols = ["part_id", "name", "source", "drawing", "material", "features", "tolerance", "batch_size", "diameter", "length", "depth",
            "width", "hole_diameter", "required_ra_um", "notes"] + (["review_status", "reviewed_by", "review_comment"] if role == "primary" else [])
    ps.append(cols)
    for p in parts:
        d = p["dims"]
        row = {"part_id": p["part_id"], "name": p["name"], "source": p["source"], "drawing": p.get("drawing", ""), "material": p["material"],
               "features": "; ".join(p["features"]), "tolerance": p["tolerance"], "batch_size": p["batch_size"], "diameter": d["diameter"],
               "length": d["length"], "depth": d["depth"], "width": d["width"], "hole_diameter": d.get("hole_diameter"),
               "required_ra_um": p.get("required_ra_um"), "notes": p.get("notes", "")}
        ps.append([row.get(c) for c in cols])
    _style_header(ps, {"A": 11, "B": 28, "C": 30, "F": 34, "O": 30})
    if role == "primary":
        dv = DataValidation(type="list", formula1='"' + ",".join(REVIEW_STATUSES) + '"', allow_blank=True)
        ps.add_data_validation(dv)
        dv.add(f"P2:P{len(parts) + 1}")

    rs = wb.create_sheet("Routes")
    head = ["part_id", "alt", "part"] + [f"step_{i}" for i in range(1, max_steps + 1)] + ["comment"]
    rs.append(head)
    for p in parts:
        for alt in ALT_LABELS[:alternatives]:
            rs.append([p["part_id"], alt, part_summary(p) if alt == "A" else None])
    _style_header(rs, {"A": 11, "B": 5, "C": 60})
    for i in range(1, max_steps + 1):
        rs.column_dimensions[openpyxl.utils.get_column_letter(3 + i)].width = 20
    rs.freeze_panes = "D2"
    last = rs.max_row
    first_step, last_step = openpyxl.utils.get_column_letter(4), openpyxl.utils.get_column_letter(3 + max_steps)
    opdv = DataValidation(type="list", formula1=f"=Operations!$A$2:$A${len(OPERATIONS) + 1}", allow_blank=True, showErrorMessage=True,
                          errorTitle="Operation", error="Pick an operation from the list (sheet 'Operations').")
    rs.add_data_validation(opdv)
    opdv.add(f"{first_step}2:{last_step}{last}")
    for row in rs.iter_rows(min_row=2, max_row=last, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    ops = wb.create_sheet("Operations")
    ops.append(["operation", "machine", "usually serves the features"])
    of = _op_features()
    for op in OPERATIONS:
        ops.append([op, machine_class(op), "; ".join(sorted(of.get(op, ())))])
    _style_header(ops, {"A": 30, "B": 10, "C": 80})

    meta = wb.create_sheet("Meta")
    for k, v in (("annotator", annotator), ("role", role), ("n_parts", len(parts)), ("alternatives", alternatives),
                 ("tool", "capp/benchmark.py (MTP-1 Step 16)")):
        meta.append([k, v])
    wb.save(path)
    return path


# ---------------------------------------------------------------- 3. import
def read_annotations(path: str) -> dict:
    """{"annotator", "role", "parts": {part_id: {"routes": {alt: [ops]}, "unknown": [(alt, text, [suggestions])], "comment",
    "review_status", "reviewed_by", "review_comment"}}} from a filled workbook. Blank cells are skipped; names are matched to the
    41 operations ignoring case and spaces."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    meta = {}
    if "Meta" in wb.sheetnames:
        for r in wb["Meta"].iter_rows(values_only=True):
            if r and r[0]:
                meta[str(r[0])] = r[1]

    def table(name):
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        header = [_text(h) for h in rows[0]]
        return header, rows[1:]

    out = {"annotator": _text(meta.get("annotator")) or "unknown", "role": _text(meta.get("role")) or "primary", "parts": {}}
    header, rows = table("Routes")
    step_cols = sorted((int(h.split("_")[1]), i) for i, h in enumerate(header) if re.fullmatch(r"step_\d+", h))
    ci = {h: i for i, h in enumerate(header)}
    for r in rows:
        pid = _text(r[ci["part_id"]])
        if not pid:
            continue
        rec = out["parts"].setdefault(pid, {"routes": {}, "unknown": [], "comment": ""})
        alt = _text(r[ci["alt"]]) or "A"
        ops = []
        for _, i in step_cols:
            cell = r[i] if i < len(r) else None
            if _blank(cell):
                continue
            op = _match(cell, OPERATIONS)
            if op is None:
                rec["unknown"].append((alt, str(cell).strip(), _suggest(cell, OPERATIONS)))
            else:
                ops.append(op)
        if ops or alt not in rec["routes"]:
            rec["routes"][alt] = ops
        if "comment" in ci and not _blank(r[ci["comment"]]):
            rec["comment"] = (rec["comment"] + " | " if rec["comment"] else "") + f"{alt}: {_text(r[ci['comment']])}"
    header, rows = table("Parts")
    ci = {h: i for i, h in enumerate(header)}
    for r in rows:
        pid = _text(r[ci["part_id"]]) if "part_id" in ci else ""
        if pid in out["parts"]:
            for k in ("review_status", "reviewed_by", "review_comment"):
                out["parts"][pid][k] = _text(r[ci[k]]) if k in ci and ci[k] < len(r) else ""
    wb.close()
    return out


def check_route(part: dict, route: list) -> dict:
    """Findings of the rule base (strict verifier) and of the physical-feasibility rules (real dimensions) for one plan."""
    chk = V.check_route(route, part["features"])
    phys = F.evaluate(route, part).violations
    return {"valid": bool(chk.valid), "kinds": list(chk.kinds), "messages": [v.message for v in chk.violations if v.kind != "machine_grouping"],
            "infeasible": sorted({v.rule_id.replace("PHYS:", "") for v in phys if v.kind == "infeasible"}),
            "risks": sorted({v.rule_id.replace("PHYS:", "") for v in phys if v.kind == "risk"})}


def _distinct(routes):
    out = []
    for r in routes:
        if r and r not in out:
            out.append(r)
    return out


def build_benchmark(spec_parts: list, primary: dict, second: dict = None, allow_missing: bool = False):
    """(bench_parts, report). `primary` and `second` come from read_annotations. Blocking problems are in report["errors"]
    (unknown operation names, a part without a plan unless allow_missing, an invalid record); report["conflicts"] lists the plans the
    rule base or the feasibility rules reject (they stay in the benchmark as references)."""
    report = {"errors": [], "missing": [], "unknown_parts": [], "conflicts": [], "physical": [], "annotator": primary["annotator"],
              "second_annotator": second["annotator"] if second else None, "review": Counter()}
    ids = {p["part_id"] for p in spec_parts}
    report["unknown_parts"] = sorted(set(primary["parts"]) - ids)
    if second:
        report["unknown_parts"] += sorted(set(second["parts"]) - ids)
    bench = []
    for p in spec_parts:
        pa = primary["parts"].get(p["part_id"])
        for who, ann in (("primary", pa), ("second", (second or {"parts": {}})["parts"].get(p["part_id"]))):
            for alt, text, sug in (ann or {}).get("unknown", []):
                report["errors"].append(f"{p['part_id']} ({who}, route {alt}): unknown operation {text!r}" + (f", did you mean {', '.join(sug)}?" if sug else ""))
        routes = _distinct([(pa or {}).get("routes", {}).get(a, []) for a in ALT_LABELS])
        if not routes or not (pa or {}).get("routes", {}).get("A"):
            report["missing"].append(p["part_id"])
            if not allow_missing:
                report["errors"].append(f"{p['part_id']}: no plan in route A of the primary workbook")
            continue
        rec = {k: v for k, v in p.items()}
        rec["reference_routes"] = routes
        meta = {"annotator": primary["annotator"], "review_status": pa.get("review_status", ""), "reviewed_by": pa.get("reviewed_by", ""),
                "review_comment": pa.get("review_comment", ""), "comment": pa.get("comment", ""), "rule_conflicts": {}, "physical": {}}
        report["review"][pa.get("review_status") or "not reviewed"] += 1
        for i, route in enumerate(routes):
            f = check_route(p, route)
            label = ALT_LABELS[i] if i < len(ALT_LABELS) else str(i)
            if not f["valid"]:
                meta["rule_conflicts"][label] = {"kinds": f["kinds"], "messages": f["messages"]}
                report["conflicts"].append({"part_id": p["part_id"], "route": label, "kinds": f["kinds"], "messages": f["messages"]})
            if f["infeasible"] or f["risks"]:
                meta["physical"][label] = {"infeasible": f["infeasible"], "risks": f["risks"]}
                report["physical"].append({"part_id": p["part_id"], "route": label, "infeasible": f["infeasible"], "risks": f["risks"]})
        if second:
            sa = second["parts"].get(p["part_id"])
            s_routes = _distinct([(sa or {}).get("routes", {}).get(a, []) for a in ALT_LABELS])
            if s_routes:
                rec["second_annotator_routes"] = s_routes
                meta["second_annotator"] = second["annotator"]
        rec["reference_meta"] = meta
        errs = D.validate_part(rec)
        report["errors"] += [f"{p['part_id']}: {e}" for e in errs]
        bench.append(rec)
    report["review"] = dict(report["review"])
    return bench, report


# ---------------------------------------------------------------- 4. agreement (M10)
def bootstrap_ci(values: list, resamples: int = BOOTSTRAP_RESAMPLES, seed: int = 7, level: float = 0.95):
    """(low, high) percentile interval of the mean; None values are ignored; (None, None) below 2 values."""
    vals = [v for v in values if v is not None]
    if len(vals) < 2:
        return None, None
    rng = random.Random(seed)
    means = sorted(statistics.fmean(rng.choices(vals, k=len(vals))) for _ in range(resamples))
    lo = means[int(((1 - level) / 2) * (resamples - 1))]
    hi = means[int((1 - (1 - level) / 2) * (resamples - 1))]
    return lo, hi


def agreement(parts: list) -> dict:
    """M10: for every part with a second plan, the M4 measures of the second person's first plan against the reference plan(s)."""
    rows = []
    for p in parts:
        s, ref = p.get("second_annotator_routes"), p.get("reference_routes")
        if s and ref:
            sim = M.best_similarity(s[0], ref)
            rows.append({"part_id": p["part_id"], "exact": sim["exact"], "lcs_ratio": sim["lcs_ratio"], "kendall_tau": sim["kendall_tau"],
                         "n_ops_reference": len(V.clean_steps(ref[0])), "n_ops_second": len(V.clean_steps(s[0]))})
    out = {"n": len(rows), "rows": rows}
    for key, vals in (("exact_rate", [1.0 if r["exact"] else 0.0 for r in rows]), ("lcs_ratio_mean", [r["lcs_ratio"] for r in rows]),
                      ("kendall_tau_mean", [r["kendall_tau"] for r in rows])):
        out[key] = M.mean_std(vals)[0]
        out[key + "_ci95"] = bootstrap_ci(vals)
    return out


# ---------------------------------------------------------------- 5. datasheet
def _f(x, nd=2):
    return "n/a" if x is None else f"{x:.{nd}f}"


def _ci(ci):
    return "n/a" if not ci or ci[0] is None else f"[{ci[0]:.2f}, {ci[1]:.2f}]"


def _table(counter: dict, head: str) -> list:
    return [f"| {head} | parts |", "|---|---|"] + [f"| {k} | {v} |" for k, v in sorted(counter.items(), key=lambda kv: (-kv[1], str(kv[0])))]


def datasheet(parts: list, report: dict = None, agree: dict = None, generated: str = "") -> str:
    """Markdown datasheet of a benchmark, computed from the records (and the import report / agreement when given)."""
    n = len(parts)
    summ = D.summarize(parts) if n else {}
    cov = D.feature_coverage(parts) if n else {"covered": 0, "total": len(fv.get_all_features()), "missing": [], "counts": {}}
    dims = {}
    for k in D.DIM_KEYS + D.OPTIONAL_DIM_KEYS:
        vals = [p["dims"][k] for p in parts if k in p["dims"]]
        dims[k] = (len(vals), min(vals), statistics.median(vals), max(vals)) if vals else (0, None, None, None)
    n_ref = Counter(len(p["reference_routes"]) for p in parts)
    lens = [len(V.clean_steps(p["reference_routes"][0])) for p in parts]
    meta = [p.get("reference_meta") or {} for p in parts]
    annotators = sorted({m.get("annotator", "") for m in meta if m.get("annotator")})
    seconds = sorted({m.get("second_annotator", "") for m in meta if m.get("second_annotator")})
    reviews = Counter((m.get("review_status") or "not reviewed") for m in meta)
    conflicts = [(p["part_id"], a, c) for p, m in zip(parts, meta) for a, c in m.get("rule_conflicts", {}).items()]
    physical = [(p["part_id"], a, c) for p, m in zip(parts, meta) for a, c in m.get("physical", {}).items()]
    L = ["# Expert benchmark: datasheet", "", f"Generated by capp/benchmark.py{(' on ' + generated) if generated else ''}. Numbers below are computed from the benchmark file; "
         "items marked [TO CONFIRM] are for the authors to fill in."]
    if any("placeholder" in (p.get("source") or "").lower() for p in parts):
        L += ["", "**WARNING: this benchmark contains synthetic PLACEHOLDER parts (not real drawings). It rehearses the pipeline; it is not the expert benchmark.**"]
    L += ["", "## Purpose",
         "Held-out evaluation of LLM process planners against plans written by people, with the human-human agreement (M10) as the ceiling "
         "for plan similarity (M4). Sealed: used for the final runs only; no prompt, rule or threshold is changed after its results are seen.", "",
         "## Composition", f"- {n} parts; {sum(1 for p in parts if p.get('source'))} with a stated source; {sum(1 for p in parts if p.get('drawing'))} with a drawing reference.",
         f"- Feature coverage: {cov['covered']}/{cov['total']} features" + (f"; not covered: {', '.join(cov['missing'])}" if cov["missing"] else "") + ".",
         f"- Real dimensions (n, min / median / max, mm): " + "; ".join(
             f"{k} {v[0]}, " + ("-" if v[0] == 0 else f"{v[1]:g} / {v[2]:g} / {v[3]:g}") for k, v in dims.items()) + ".",
         f"- Required Ra given for {sum(1 for p in parts if p.get('required_ra_um'))} parts; hole diameter given for {dims['hole_diameter'][0]} parts.", ""]
    if n:
        L += _table(Counter(p.get("source", "") or "(none)" for p in parts), "source") + [""]
        L += _table(summ["materials"], "material") + [""] + _table(summ["tolerances"], "tolerance") + [""]
        L += _table(summ["machine_profiles"], "machine profile") + [""] + _table(summ["batch_categories"], "batch category") + [""]
        L += _table({f"{k} features": v for k, v in summ["n_features_hist"].items()}, "features per part") + [""]
    L += ["## Reference plans",
          f"- Primary annotator: {', '.join(annotators) or 'unknown'}. Supervisor review: " + ", ".join(f"{k} {v}" for k, v in sorted(reviews.items())) + ".",
          f"- Plans per part (alternatives the annotator accepts): " + ", ".join(f"{k}: {v} parts" for k, v in sorted(n_ref.items())) + ".",
          f"- Operations in the first plan: mean {_f(statistics.fmean(lens) if lens else None, 1)}, min {min(lens) if lens else 'n/a'}, max {max(lens) if lens else 'n/a'}.",
          f"- Plans the rule base (precedence, missing / invented operations) rejects: {len(conflicts)}"
          + (" (" + ", ".join(f"{pid}/{a}: {'+'.join(c['kinds'])}" for pid, a, c in conflicts[:12]) + (", ..." if len(conflicts) > 12 else "") + ")" if conflicts else "")
          + ". They stay as references: the expert is the ground truth, and each conflict is either an expert slip or a gap of the rule base, to be settled with the supervisor.",
          f"- Plans with a physical-feasibility finding (real dimensions): {len(physical)}"
          + (" (" + ", ".join(f"{pid}/{a}: {'+'.join(c['infeasible'] + c['risks'])}" for pid, a, c in physical[:12]) + (", ..." if len(physical) > 12 else "") + ")" if physical else "") + ".", ""]
    if report:
        L += [f"- Import: {len(report.get('missing', []))} spec part(s) without a plan (excluded), {len(report.get('errors', []))} blocking error(s) at import.", ""]
    L += ["## Human-human agreement (M10)"]
    if agree and agree["n"]:
        L += [f"- Second annotator: {', '.join(seconds) or 'unknown'}; {agree['n']} parts, plan written independently from the same specification.",
              f"- Exact match with the reference: {_f(agree['exact_rate'])} {_ci(agree['exact_rate_ci95'])}; LCS ratio {_f(agree['lcs_ratio_mean'])} {_ci(agree['lcs_ratio_mean_ci95'])}; "
              f"Kendall tau distance {_f(agree['kendall_tau_mean'])} {_ci(agree['kendall_tau_mean_ci95'])} (0 = same order). Intervals: 95% bootstrap over parts.",
              "- Read as a ceiling for M4 on this benchmark, with the width of those intervals: with 15 to 20 parts it is indicative, not precise.", ""]
    else:
        L += ["- No second-annotator plans in this benchmark yet. M10 is not available.", ""]
    L += ["## Intended use and limits",
          "- Final runs only (WP11). Report M4 together with M10, and the number of parts.",
          "- Reference plans reflect one expert's practice (the alternatives recorded per part and the supervisor review mitigate this, they do not remove it).",
          "- The benchmark is small; per-class tables on it are anecdotes. The rule base's 'Inspection last' convention is part of what the plans were asked to follow.",
          "- Dimensions are the envelope of the part, not a full drawing; the process plan ignores fixtures and tolerances finer than the 6 categories.", "",
          "## To fill in by hand", "- [TO CONFIRM] sources and licences of each drawing or textbook example.", "- [TO CONFIRM] annotators' background and years of experience.",
          "- [TO CONFIRM] date of supervisor sign-off and the tag of the repository at which the benchmark was frozen."]
    return "\n".join(L) + "\n"
